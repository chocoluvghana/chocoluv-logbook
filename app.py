from datetime import date
import hashlib
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import plotly.express as px
from sqlalchemy import URL, create_engine, text
import streamlit as st
import streamlit.components.v1 as components
import re

# --- Page Configuration ---
st.set_page_config(
    page_title="CHOCOLUV LTD SALES LOGBOOK",
    page_icon="🍫",
    layout="wide",
    initial_sidebar_state="expanded",
)

# --- Dual-Database Connection & Auto-Table Initialization ---
try:
  if "DATABASE_URL" in st.secrets:
    DB_URI = st.secrets["DATABASE_URL"]
  else:
    DB_URI = "postgresql+psycopg://neondb_owner:npg_p4rhxY6HwWCL@ep-aged-firefly-b5b01bzp-pooler.c-7.us-east-2.aws.neon.tech/neondb?sslmode=require"
except Exception:
  DB_URI = "postgresql+psycopg://neondb_owner:npg_p4rhxY6HwWCL@ep-aged-firefly-b5b01bzp-pooler.c-7.us-east-2.aws.neon.tech/neondb?sslmode=require"

local_pass = st.secrets.get("LOCAL_DB_PASSWORD", "Chocoluv@2026")

LOCAL_DB_URI = URL.create(
    drivername="postgresql+psycopg",
    username="postgres",
    password=local_pass,
    host="127.0.0.1",
    port=5432,
    database="chocoluv_db",
)


@st.cache_resource
def get_db_engines():
  engine_neon = create_engine(
      DB_URI,
      pool_pre_ping=True,
      pool_recycle=300,
      connect_args={"connect_timeout": 15},
  )
  
  init_sql = """
      CREATE TABLE IF NOT EXISTS local_sales (
          id SERIAL PRIMARY KEY,
          date DATE,
          customer_name TEXT,
          batch_number TEXT,
          product_name TEXT,
          weight_g FLOAT,
          quantity INT,
          unit_cost FLOAT,
          exchange_rate FLOAT DEFAULT 1.0,
          entered_by TEXT
      );
      CREATE TABLE IF NOT EXISTS export_sales (
          id SERIAL PRIMARY KEY,
          date DATE,
          customer_name TEXT,
          batch_number TEXT,
          product_name TEXT,
          weight_g FLOAT,
          quantity INT,
          unit_cost FLOAT,
          exchange_rate FLOAT DEFAULT 1.0,
          entered_by TEXT
      );
      CREATE TABLE IF NOT EXISTS protocol_sales (
          id SERIAL PRIMARY KEY,
          date DATE,
          customer_name TEXT,
          batch_number TEXT,
          product_name TEXT,
          weight_g FLOAT,
          quantity INT,
          unit_cost FLOAT,
          exchange_rate FLOAT DEFAULT 1.0,
          entered_by TEXT
      );
      CREATE TABLE IF NOT EXISTS staff_users (
          username TEXT PRIMARY KEY,
          password_hash TEXT
      );
  """

  with engine_neon.begin() as conn:
    conn.execute(text(init_sql))

  engine_local = None
  try:
    engine_local = create_engine(
        LOCAL_DB_URI,
        pool_pre_ping=True,
        pool_recycle=300,
        connect_args={"connect_timeout": 2},
    )
    with engine_local.begin() as conn:
      conn.execute(text(init_sql))
  except Exception:
    engine_local = None

  return engine_neon, engine_local


engine, engine_local = get_db_engines()

TABLE_MAP = {
    "Local Sales": "local_sales",
    "Exports": "export_sales",
    "Protocol": "protocol_sales",
}

# --- Load Product Catalog from Excel with Embedded Fallback ---
@st.cache_data
def load_product_catalog():
  try:
    df_cat = pd.read_excel("chocoluv_product_catalog.xlsx", sheet_name="Sheet1")
    products = ["-- Select Product --"] + df_cat["Product Name"].dropna().tolist()
    return products, df_cat
  except Exception:
    data = [
        ("Best Wishes (41g)", "41g", 25.0, 25.0),
        ("Bunny / Barney (36g)", "36g", 25.0, 25.0),
        ("SpongeBob (36g)", "36g", 25.0, 25.0),
        ("Spelt Happy Birthday (36g)", "36g", 25.0, 25.0),
        ("Parcel / Present Happy Birthday (36g)", "36g", 25.0, 25.0),
        ("Train (69g)(Big)", "69g", 30.0, 30.0),
        ("Aeroplane (61g)(Big)", "61g", 30.0, 30.0),
        ("Bear Feet / Foot Print (43g)", "43g", 25.0, 25.0),
        ("Round Happy Birthday", "45g", 25.0, 25.0),
        ("Super Star", "40g", 25.0, 25.0),
        ("Flower Bouquet(Big)", "50g", 30.0, 30.0),
        ("Ninja Turtle", "36g", 25.0, 25.0),
        ("Choco Swirl", "25g", 25.0, 25.0),
        ("6g Milk Square Nugget", "6g", 6.0, 7.0),
        ("12g Milk Square Nugget", "12g", 8.0, 8.0),
        ("12g Milk Chocolate Bar", "12g", 8.0, 8.0),
        ("12g Dark Chocolate Bar", "12g", 8.0, 8.0),
        ("27g Milk 50%", "27g", 12.0, 25.0),
        ("27g Dark 72%", "27g", 12.0, 25.0),
        ("27g Milk Almond", "27g", 12.0, 30.0),
        ("50g Milk 50%", "50g", 16.0, 50.0),
        ("50g Dark 72%", "50g", 16.0, 50.0),
        ("50g 85% Dark Bar", "50g", 16.0, 70.0),
        ("50g Dark with Fruit / Orange", "50g", 16.0, 50.0),
        ("50g Dark Coffee 72%", "50g", 16.0, 50.0),
        ("50g Dark Ginger 72%", "50g", 16.0, 50.0),
        ("50g Dark Nibs 72%", "50g", 16.0, 50.0),
        ("50g Dark with Almond", "50g", 16.0, 60.0),
        ("50g Dark with Kelewele", "50g", 16.0, 50.0),
        ("50g Dark with Cashew/Millet", "50g", 16.0, 50.0),
        ("50g Dark with Chilli", "50g", 16.0, 50.0),
        ("50g Dark Calabash Nutmeg", "50g", 16.0, 50.0),
        ("50g Milk Coffee 50%", "50g", 16.0, 50.0),
        ("50g Milk Ginger 50%", "50g", 16.0, 50.0),
        ("50g Milk Orange 50%", "50g", 16.0, 50.0),
        ("50g Milk with Almond", "50g", 16.0, 60.0),
        ("50g Milk with Kelewele", "50g", 16.0, 50.0),
        ("50g Milk with Calabash Nutmeg", "50g", 16.0, 50.0),
        ("50g Milk with Cashew/Millet", "50g", 16.0, 50.0),
        ("100g Milk Bar", "100g", 26.0, 100.0),
        ("100g Dark Bar", "100g", 26.0, 100.0),
        ("50g 85% Dark Sugar Free", "50g", 16.0, 70.0),
        ("12 Piece Dark Box (125g)", "125g", 50.0, 150.0),
        ("12 Piece Milk Box (125g)", "125g", 50.0, 150.0),
        ("12 Piece Assorted Box (125g)", "125g", 50.0, 150.0),
        ("8 Piece Chocolate Pack (27g)", "216g", 40.0, 80.0),
        ("20 Piece Chocolate Pack (27g)", "540g", 161.0, 500.0),
        ("6 Piece Assorted Box", "300g", 83.0, 300.0),
        ("8 Piece Assorted Box", "400g", 100.0, 350.0),
        ("10 Piece Assorted Box", "500g", 125.0, 250.0),
        ("20 Piece Assorted Box", "1000g", 161.0, 500.0),
        ("Assorted Box (12 piece)", "125g", 50.0, 150.0),
        ("Dark Box (12 piece)", "125g", 50.0, 150.0),
        ("Milk Box (12 piece)", "125g", 50.0, 150.0),
        ("Oatmeal (Plain)", "44g", 10.0, 15.0),
        ("Oatmeal (Chocolate)", "70g", 12.0, 20.0),
        ("Mango Balls", "100g", 100.0, 100.0),
        ("Pineapple Dip", "100g", 50.0, 100.0),
        ("Pawpaw Dip", "100g", 50.0, 100.0),
        ("Mixed Fruit Dip", "100g", 50.0, 100.0),
        ("Instant Drinking Chocolate", "250g", 61.0, 120.0),
        ("Natural Cocoa Powder", "250g", 68.0, 150.0),
        ("Raw Nibs", "100g", 100.0, 100.0),
        ("Mango Dip", "100g", 50.0, 100.0),
        ("Fruit balls", "100g", 100.0, 100.0),
        ("Heart Nuggets", "10g", 5.0, 5.0),
    ]
    df_cat = pd.DataFrame(
        data,
        columns=["Product Name", "Weight", "Local Price (GHS)", "Export Price (GHS)"],
    )
    return ["-- Select Product --"] + df_cat["Product Name"].tolist(), df_cat


PRODUCT_CATALOG, CATALOG_DF = load_product_catalog()


def parse_weight_to_grams(w_val, product_name=""):
  if product_name and "almond" in product_name.lower():
    if "27g" not in product_name.lower():
      return 80.0

  if pd.isna(w_val):
    return 50.0
  w_str = str(w_val).strip()
  match = re.search(r"(\d+(?:\.\d+)?)", w_str)
  if match:
    val = float(match.group(1))
    return val * 1000.0 if "kg" in w_str.lower() else val
  return 50.0


# --- Database Operations ---
def fetch_table(table_name):
  active_engine = engine if engine_local is None else engine_local
  try:
    query = f"SELECT * FROM {table_name} ORDER BY date DESC, id DESC"
    df = pd.read_sql(query, active_engine)
  except Exception:
    try:
      query_fallback = f'SELECT * FROM "{table_name}"'
      df = pd.read_sql(query_fallback, active_engine)
    except Exception:
      df = pd.DataFrame()

  if not df.empty:
    if "date" in df.columns:
      df["date"] = pd.to_datetime(df["date"]).dt.date
    if "weight_g" in df.columns and "quantity" in df.columns:
      df["total_weight_g"] = df["quantity"] * df["weight_g"]
      df["total_weight_mt"] = df["total_weight_g"] / 1_000_000.0
    if "unit_cost" in df.columns and "quantity" in df.columns:
      df["total_amount"] = df["quantity"] * df["unit_cost"]
    if "exchange_rate" not in df.columns:
      df["exchange_rate"] = 1.0000
    if "batch_number" not in df.columns:
      df["batch_number"] = "-"
  return df


def insert_dispatch(table_name, record):
  insert_sql = text(f"""
        INSERT INTO {table_name} (date, customer_name, batch_number, product_name, weight_g, quantity, unit_cost, entered_by)
        VALUES (:date, :customer_name, :batch_number, :product_name, :weight_g, :quantity, :unit_cost, :entered_by)
    """)
  with engine.begin() as conn:
    conn.execute(insert_sql, record)
  if engine_local:
    try:
      with engine_local.begin() as conn_loc:
        conn_loc.execute(insert_sql, record)
    except Exception:
      pass


def update_dispatch_record(table_name, record_id, record):
  update_sql = text(f"""
        UPDATE {table_name}
        SET date = :date,
            customer_name = :customer_name,
            batch_number = :batch_number,
            product_name = :product_name,
            weight_g = :weight_g,
            quantity = :quantity,
            unit_cost = :unit_cost,
            entered_by = :entered_by
        WHERE id = :id
    """)
  record["id"] = record_id
  with engine.begin() as conn:
    conn.execute(update_sql, record)
  if engine_local:
    try:
      with engine_local.begin() as conn_loc:
        conn_loc.execute(update_sql, record)
    except Exception:
      pass


def delete_record(table_name, record_id):
  sql = text(f"DELETE FROM {table_name} WHERE id = :id")
  with engine.begin() as conn:
    conn.execute(sql, {"id": record_id})
  if engine_local:
    try:
      with engine_local.begin() as conn_loc:
        conn_loc.execute(sql, {"id": record_id})
    except Exception:
      pass


def apply_monthly_rate(table_name, year, month, rate):
  sql = text(f"""
        UPDATE {table_name}
        SET exchange_rate = :rate
        WHERE EXTRACT(YEAR FROM date) = :yr AND EXTRACT(MONTH FROM date) = :mo
    """)
  with engine.begin() as conn:
    conn.execute(sql, {"rate": rate, "yr": year, "mo": month})
  if engine_local:
    try:
      with engine_local.begin() as conn_loc:
        conn_loc.execute(sql, {"rate": rate, "yr": year, "mo": month})
    except Exception:
      pass


def update_row_rate(table_name, record_id, rate):
  sql = text(f"""
        UPDATE {table_name}
        SET exchange_rate = :rate
        WHERE id = :id
    """)
  with engine.begin() as conn:
    conn.execute(sql, {"rate": rate, "id": record_id})
  if engine_local:
    try:
      with engine_local.begin() as conn_loc:
        conn_loc.execute(sql, {"rate": rate, "id": record_id})
    except Exception:
      pass


# --- Staff Database Authentication Functions ---
def hash_text(secret: str) -> str:
  return hashlib.sha256(secret.strip().encode("utf-8")).hexdigest()


def get_staff_account(username: str):
  active_engine = engine if engine_local is None else engine_local
  sql = text(
      "SELECT username, password_hash FROM staff_users WHERE"
      " LOWER(username)=LOWER(:u) LIMIT 1"
  )
  with active_engine.connect() as conn:
    res = conn.execute(sql, {"u": username.strip()}).fetchone()
    return res


def register_staff_account(username: str, secret: str):
  sql = text("""
        INSERT INTO staff_users (username, password_hash)
        VALUES (:u, :p)
    """)
  with engine.begin() as conn:
    conn.execute(sql, {"u": username.strip().title(), "p": hash_text(secret)})
  if engine_local:
    try:
      with engine_local.begin() as conn_loc:
        conn_loc.execute(sql, {"u": username.strip().title(), "p": hash_text(secret)})
    except Exception:
      pass


# --- Universal Theme Immunity ---
st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=Cinzel:wght@700&display=swap');

.stApp, [data-testid="stAppViewContainer"] {
    background-color: #140a06 !important;
    background-image: 
        radial-gradient(circle at 10% 15%, rgba(162, 92, 53, 0.40) 0%, transparent 45%),
        radial-gradient(circle at 90% 85%, rgba(70, 36, 20, 0.55) 0%, transparent 50%),
        linear-gradient(150deg, #180d08 0%, #0d0604 100%) !important;
    background-attachment: fixed !important;
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    color: #ffffff !important;
}

section[data-testid="stSidebar"], 
div[data-testid="stSidebarNav"],
[data-testid="stSidebar"] > div:first-child {
    background-color: #1f110a !important;
    background-image: linear-gradient(180deg, #24140c 0%, #160c07 100%) !important;
    border-right: 1.5px solid #d4a373 !important;
}
section[data-testid="stSidebar"] * {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
}
section[data-testid="stSidebar"] button {
    background: #d48b50 !important;
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    border: none !important;
    border-radius: 8px !important;
    font-weight: 700 !important;
}

[data-testid="stForm"], 
div[data-testid="stMetric"], 
.stTabs [data-baseweb="tab-panel"] {
    background-color: rgba(36, 20, 14, 0.92) !important;
    backdrop-filter: blur(14px) !important;
    -webkit-backdrop-filter: blur(14px) !important;
    border: 1.5px solid rgba(212, 163, 115, 0.45) !important;
    border-radius: 16px !important;
    padding: 22px !important;
    box-shadow: 0 10px 30px rgba(0, 0, 0, 0.6) !important;
}

h1, h2, h3, h4, 
[data-testid="stForm"] h3, 
.stTabs [data-baseweb="tab-panel"] h3 {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    font-weight: 800 !important;
}
.header-badge {
    background: rgba(212, 163, 115, 0.2);
    color: #f4a261;
    padding: 6px 16px;
    border-radius: 20px;
    font-size: 0.82rem;
    font-weight: 800;
    letter-spacing: 1.5px;
    text-transform: uppercase;
    border: 1.5px solid #d4a373;
    display: inline-block;
    margin-bottom: 8px;
}
.header-title {
    font-family: 'Cinzel', serif !important;
    font-size: 2.6rem !important;
    font-weight: 800 !important;
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    margin: 0 !important;
}
.header-desc {
    color: #eddcd2 !important;
    -webkit-text-fill-color: #eddcd2 !important;
    font-size: 1rem !important;
    margin-top: 4px !important;
    margin-bottom: 22px !important;
}

[data-testid="stWidgetLabel"] p, 
[data-testid="stAppViewContainer"] label p,
[data-testid="stMarkdownContainer"] p {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    font-size: 0.95rem !important;
    font-weight: 700 !important;
}

div[data-baseweb="input"],
div[data-baseweb="base-input"],
div[data-baseweb="select"] > div,
div[data-testid="stTextInput"] input,
div[data-testid="stNumberInput"] input,
div[data-testid="stDateInput"] input {
    background-color: #ffffff !important;
    border: 1.5px solid #c9935d !important;
    border-radius: 9px !important;
    color: #000000 !important;
    -webkit-text-fill-color: #000000 !important;
    font-weight: 800 !important;
    font-size: 0.98rem !important;
}

div[data-baseweb="select"] span,
div[data-baseweb="select"] div,
div[data-baseweb="select"] svg {
    color: #000000 !important;
    -webkit-text-fill-color: #000000 !important;
    fill: #000000 !important;
    font-weight: 800 !important;
}

::placeholder,
input::placeholder {
    color: #555555 !important;
    -webkit-text-fill-color: #555555 !important;
    font-weight: 600 !important;
}

button[kind="primary"] {
    background: linear-gradient(135deg, #d48b50 0%, #873e23 100%) !important;
    border: 1px solid #f4a261 !important;
    border-radius: 10px !important;
    font-weight: 800 !important;
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
    padding: 10px 22px !important;
    box-shadow: 0 4px 15px rgba(0,0,0,0.4) !important;
}
</style>
""",
    unsafe_allow_html=True,
)


def check_password():
  if st.session_state.get("authenticated", False):
    return True

  st.markdown("### 🔒 Chocoluv Staff Portal")
  auth_mode = st.radio(
      "Select Action",
      ["Log In", "First Time? Set Up Your Password"],
      horizontal=True,
  )

  if auth_mode == "Log In":
    with st.form("login_form"):
      login_user = st.text_input("Your Name / Username")
      login_pass = st.text_input("Password", type="password")
      if st.form_submit_button("Log In", type="primary", use_container_width=True):
        if not login_user.strip() or not login_pass.strip():
          st.error("Please enter both your name and password.")
        else:
          try:
            account = get_staff_account(login_user)
            if account and account[1] == hash_text(login_pass):
              st.session_state["authenticated"] = True
              st.session_state["staff_user"] = account[0]
              st.success(f"Welcome back, {account[0]}!")
              st.rerun()
            else:
              st.error("Incorrect username or password.")
          except Exception:
            st.error("Database connection error. Please try again.")
  else:
    with st.form("register_form"):
      new_user = st.text_input("Your Name (e.g. Angela)")
      new_pass = st.text_input("Choose a Password", type="password")
      confirm_pass = st.text_input("Confirm Password", type="password")
      if st.form_submit_button("Save & Register My Account", type="primary", use_container_width=True):
        clean_user = new_user.strip().title()
        if not clean_user or not new_pass.strip():
          st.error("Please fill in both your name and password.")
        elif new_pass != confirm_pass:
          st.error("Passwords do not match.")
        else:
          try:
            if get_staff_account(clean_user):
              st.error("An account for this name already exists!")
            else:
              register_staff_account(clean_user, new_pass)
              st.success("Account created successfully!")
              st.session_state["authenticated"] = True
              st.session_state["staff_user"] = clean_user
              st.rerun()
          except Exception:
            st.error("Database connection error. Please try again.")
  return False


st.markdown(
    """
<div>
    <span class="header-badge">✦ Sales & QA Operations Console</span>
    <h1 class="header-title">CHOCOLUV LTD SALES LOGBOOK</h1>
    <p class="header-desc">Record customer orders, manage batch numbers, exchange rates, and audit financial records.</p>
</div>
""",
    unsafe_allow_html=True,
)

if not check_password():
  st.stop()

with st.sidebar:
  st.markdown(f"👤 Logged in as: **{st.session_state.get('staff_user')}**")
  if st.button("Log Out", use_container_width=True):
    st.session_state["authenticated"] = False
    st.session_state["staff_user"] = ""
    st.rerun()


def to_excel_bytes(df):
  if df.empty:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
      df.to_excel(writer, index=False)
    return output.getvalue()

  work_df = df.copy()
  work_df["date"] = pd.to_datetime(work_df["date"]).dt.date
  if "exchange_rate" not in work_df.columns:
    work_df["exchange_rate"] = 1.0000
  if "batch_number" not in work_df.columns:
    work_df["batch_number"] = "-"

  work_df["batch_number"] = work_df["batch_number"].fillna("-")
  work_df["item_total_weight_g"] = work_df["quantity"] * work_df["weight_g"]
  work_df["item_total_amount_ghc"] = work_df["quantity"] * work_df["unit_cost"]

  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Sales Log"
  ws.views.sheetView[0].showGridLines = True

  header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
  header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
  thin_border = Border(
      left=Side(style="thin", color="DDDDDD"),
      right=Side(style="thin", color="DDDDDD"),
      top=Side(style="thin", color="DDDDDD"),
      bottom=Side(style="thin", color="DDDDDD"),
  )

  headers = [
      "ID", "DATE", "CLIENT", "BATCH #", "PRODUCT", "WEIGHT (g)", "QTY",
      "UNIT COST GH₵", "TOTAL AMOUNT GH₵", "TOTAL WEIGHT (g)",
      "TOTAL WEIGHT IN M/T", "TOTAL AMOUNT OF ORDER (GH₵)",
      "EXCHANGE RATE", "TOTAL AMOUNT IN USD ($)"
  ]

  ws.row_dimensions[1].height = 28
  for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

  current_row = 2
  grouped = work_df.groupby(["date", "customer_name"], sort=False)

  for (d_val, client), group in grouped:
    num_rows = len(group)
    start_row = current_row
    end_row = current_row + num_rows - 1

    order_total_weight_g = group["item_total_weight_g"].sum()
    order_total_weight_mt = order_total_weight_g / 1_000_000.0
    order_total_ghc = group["item_total_amount_ghc"].sum()
    order_rate = group["exchange_rate"].iloc[0] if "exchange_rate" in group.columns else 11.20
    order_total_usd = round(order_total_ghc / order_rate, 2) if order_rate > 0 else 0.00

    for idx, (_, row) in enumerate(group.iterrows()):
      r = current_row
      ws.row_dimensions[r].height = 22
      ws.cell(row=r, column=1, value=f"#{row['id']}").alignment = Alignment(horizontal="center", vertical="center")
      ws.cell(row=r, column=4, value=str(row["batch_number"]) if pd.notna(row["batch_number"]) else "-").alignment = Alignment(horizontal="center", vertical="center")
      ws.cell(row=r, column=5, value=str(row["product_name"])).alignment = Alignment(horizontal="left", vertical="center")

      ws.cell(row=r, column=6, value=float(row["weight_g"])).number_format = "#,##0.0"
      ws.cell(row=r, column=7, value=int(row["quantity"])).number_format = "#,##0"
      ws.cell(row=r, column=8, value=float(row["unit_cost"])).number_format = "#,##0.00"
      ws.cell(row=r, column=9, value=float(row["item_total_amount_ghc"])).number_format = "#,##0.00"

      for c_idx in range(1, 15):
        ws.cell(row=r, column=c_idx).border = thin_border
      current_row += 1

    ws.cell(row=start_row, column=2, value=str(d_val)).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=start_row, column=3, value=str(client)).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=start_row, column=10, value=float(order_total_weight_g)).number_format = "#,##0.00"
    ws.cell(row=start_row, column=11, value=float(order_total_weight_mt)).number_format = "#,##0.000000"
    ws.cell(row=start_row, column=12, value=float(order_total_ghc)).number_format = '"GH₵ "#,##0.00'
    ws.cell(row=start_row, column=13, value=float(order_rate)).number_format = "0.0000"
    ws.cell(row=start_row, column=14, value=float(order_total_usd)).number_format = '"$"#,##0.00'

    if num_rows > 1:
      for col in [2, 3, 10, 11, 12, 13, 14]:
        ws.merge_cells(start_row=start_row, start_column=col, end_row=end_row, end_column=col)

  output = io.BytesIO()
  wb.save(output)
  return output.getvalue()


def display_merged_excel_table(df):
  if df.empty:
    st.info("No records to display.")
    return

  work_df = df.copy()
  work_df["date"] = pd.to_datetime(work_df["date"]).dt.date
  if "exchange_rate" not in work_df.columns:
    work_df["exchange_rate"] = 1.0000
  if "batch_number" not in work_df.columns:
    work_df["batch_number"] = "-"

  work_df["batch_number"] = work_df["batch_number"].fillna("-")
  work_df["item_total_weight_g"] = work_df["quantity"] * work_df["weight_g"]
  work_df["item_total_amount_ghc"] = work_df["quantity"] * work_df["unit_cost"]

  html_parts = ["""
    <!DOCTYPE html>
    <html>
    <head>
    <style>
        body { margin: 0; padding: 0; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; background: transparent; }
        .table-wrap { width: 100%; overflow-x: auto; border-radius: 12px; border: 1px solid #d4a373; box-shadow: 0 4px 16px rgba(0,0,0,0.3); }
        table { width: 100%; border-collapse: collapse; background-color: #ffffff; color: #1a0f0b; font-size: 13px; text-align: center; }
        thead tr { background: #2e7d32; color: #ffffff; font-weight: 800; font-size: 12px; letter-spacing: 0.5px; }
        th { padding: 12px 8px; border: 1px solid #c8e6c9; white-space: nowrap; }
        td { padding: 9px 8px; border: 1px solid #e0e0e0; vertical-align: middle; }
        .id-cell { font-weight: 700; color: #888888; font-size: 11px; background: #fafafa; }
        .client-cell { font-weight: 800; color: #6b3310; background: #fffcf8; text-transform: uppercase; }
        .date-cell { font-weight: 700; background: #fafafa; }
        .batch-cell { font-weight: 700; color: #2e7d32; font-family: monospace; background: #f9fbf9; }
        .prod-cell { text-align: left; font-weight: 600; }
        .qty-cell { font-weight: 700; }
        .order-sum { font-weight: 800; background: #fdfaf6; }
        .order-usd { font-weight: 800; color: #1b5e20; background: #f2fbf2; }
        tr:hover { background-color: #faf4ee; }
    </style>
    </head>
    <body>
    <div class="table-wrap">
    <table>
        <thead>
            <tr>
                <th>ID</th>
                <th>DATE</th>
                <th>CLIENT</th>
                <th>BATCH #</th>
                <th style="text-align:left;">PRODUCT</th>
                <th>WEIGHT (g)</th>
                <th>QTY</th>
                <th>UNIT COST GH₵</th>
                <th>TOTAL AMOUNT GH₵</th>
                <th>TOTAL WEIGHT (g)</th>
                <th>TOTAL WEIGHT IN M/T</th>
                <th>TOTAL AMOUNT OF ORDER (GH₵)</th>
                <th>EXCHANGE RATE</th>
                <th>TOTAL AMOUNT IN USD ($)</th>
            </tr>
        </thead>
        <tbody>
    """]

  grouped = work_df.groupby(["date", "customer_name"], sort=False)
  total_rendered_rows = 0
  for (d_val, client), group in grouped:
    num_rows = len(group)
    total_rendered_rows += num_rows

    order_total_weight_g = group["item_total_weight_g"].sum()
    order_total_weight_mt = order_total_weight_g / 1_000_000.0
    order_total_ghc = group["item_total_amount_ghc"].sum()
    order_rate = group["exchange_rate"].iloc[0] if "exchange_rate" in group.columns else 11.20
    order_total_usd = round(order_total_ghc / order_rate, 2) if order_rate > 0 else 0.00

    for idx, (_, row) in enumerate(group.iterrows()):
      html_parts.append("<tr>")
      html_parts.append(f"<td class='id-cell'>#{row['id']}</td>")
      if idx == 0:
        html_parts.append(f"<td class='date-cell' rowspan='{num_rows}'>{d_val}</td>")
        html_parts.append(f"<td class='client-cell' rowspan='{num_rows}'>{client}</td>")

      b_val = str(row["batch_number"]) if pd.notna(row["batch_number"]) else "-"
      html_parts.append(f"<td class='batch-cell'>{b_val}</td>")
      html_parts.append(f"<td class='prod-cell'>{row['product_name']}</td>")
      html_parts.append(f"<td>{row['weight_g']:,.1f}</td>")
      html_parts.append(f"<td class='qty-cell'>{row['quantity']}</td>")
      html_parts.append(f"<td>{row['unit_cost']:,.2f}</td>")
      html_parts.append(f"<td style='font-weight:700;'>{row['item_total_amount_ghc']:,.2f}</td>")

      if idx == 0:
        html_parts.append(f"<td class='order-sum' rowspan='{num_rows}'>{order_total_weight_g:,.2f}</td>")
        html_parts.append(f"<td class='order-sum' rowspan='{num_rows}'>{order_total_weight_mt:,.6f}</td>")
        html_parts.append(f"<td class='order-sum' rowspan='{num_rows}' style='color:#b25e24;'>GH₵ {order_total_ghc:,.2f}</td>")
        html_parts.append(f"<td rowspan='{num_rows}' style='font-weight:700;'>{order_rate:.4f}</td>")
        html_parts.append(f"<td class='order-usd' rowspan='{num_rows}'>${order_total_usd:,.2f}</td>")

      html_parts.append("</tr>")

  html_parts.append("</tbody></table></div></body></html>")
  full_table_html = "".join(html_parts)
  calculated_height = max(180, (total_rendered_rows * 42) + 70)
  components.html(full_table_html, height=calculated_height, scrolling=True)


def prepare_export_df(df):
  export_df = df.copy()
  export_df["date"] = pd.to_datetime(export_df["date"]).dt.date
  if "exchange_rate" not in export_df.columns:
    export_df["exchange_rate"] = 1.0000
  if "batch_number" not in export_df.columns:
    export_df["batch_number"] = "-"

  export_df["batch_number"] = export_df["batch_number"].fillna("-")
  export_df["item_total_weight_g"] = export_df["quantity"] * export_df["weight_g"]
  export_df["item_total_amount_ghc"] = export_df["quantity"] * export_df["unit_cost"]

  order_weights = export_df.groupby(["date", "customer_name"])["item_total_weight_g"].transform("sum")
  order_totals_ghc = export_df.groupby(["date", "customer_name"])["item_total_amount_ghc"].transform("sum")

  export_df["TOTAL WEIGHT (g)"] = order_weights
  export_df["TOTAL WEIGHT IN M/T"] = (order_weights / 1_000_000.0).round(6)
  export_df["TOTAL AMOUNT OF ORDER (GH₵)"] = order_totals_ghc
  export_df["TOTAL AMOUNT IN USD ($)"] = (order_totals_ghc / export_df["exchange_rate"]).round(2)

  rename_map = {
      "id": "RECORD ID",
      "date": "DATE",
      "customer_name": "CLIENT",
      "batch_number": "BATCH NUMBER",
      "product_name": "PRODUCT",
      "weight_g": "WEIGHT (g)",
      "quantity": "QTY",
      "unit_cost": "UNIT COST (GH₵)",
      "item_total_amount_ghc": "TOTAL AMOUNT GH₵",
      "exchange_rate": "EXCHANGE RATE",
      "entered_by": "ENTERED BY",
  }
  return export_df.rename(columns=rename_map)


def render_table_edit_delete_bar(table_name, key_prefix):
  with st.expander("✏ Quick Edit / Delete a Record", expanded=False):
    current_df = fetch_table(table_name)
    if current_df.empty:
      st.info("No records in this table to edit or delete.")
      return

    if "batch_number" not in current_df.columns:
      current_df["batch_number"] = ""

    current_df["display_label"] = current_df.apply(
        lambda r: f"Record #{r['id']} | {r['date']} | {r['customer_name']} | Product: {r['product_name']} (Qty: {r['quantity']})",
        axis=1,
    )

    selected_label = st.selectbox(
        "Select Record to Edit or Delete:",
        current_df["display_label"].tolist(),
        key=f"{key_prefix}_sel",
    )
    selected_row = current_df[current_df["display_label"] == selected_label].iloc[0]
    rec_id = int(selected_row["id"])

    action_mode = st.radio("Choose Action", ["Edit Record", "Delete Record"], horizontal=True, key=f"{key_prefix}_mode")

    if action_mode == "Edit Record":
      with st.form(f"{key_prefix}_edit_form"):
        col1, col2, col3 = st.columns([2, 2, 2])
        with col1:
          e_date = st.date_input("Date", value=pd.to_datetime(selected_row["date"]).date(), key=f"{key_prefix}_d")
        with col2:
          e_cust = st.text_input("Customer Name", value=selected_row["customer_name"], key=f"{key_prefix}_c")
        with col3:
          curr_batch = selected_row["batch_number"] if pd.notna(selected_row["batch_number"]) else ""
          e_batch = st.text_input("Batch Number", value=curr_batch, key=f"{key_prefix}_b")

        col4, col5, col6 = st.columns([3, 1, 2])
        with col4:
          prod_idx = PRODUCT_CATALOG.index(selected_row["product_name"]) if selected_row["product_name"] in PRODUCT_CATALOG else 0
          e_prod = st.selectbox("Product", PRODUCT_CATALOG, index=prod_idx, key=f"{key_prefix}_p")
        with col5:
          e_weight = st.number_input("Weight (g)", min_value=0.0, step=1.0, value=float(selected_row["weight_g"]), format="%.2f", key=f"{key_prefix}_w")
        with col6:
          e_staff = st.text_input("Entered By", value=selected_row["entered_by"], key=f"{key_prefix}_s")

        col7, col8 = st.columns(2)
        with col7:
          e_qty = st.number_input("Quantity", min_value=1, step=1, value=int(selected_row["quantity"]), key=f"{key_prefix}_q")
        with col8:
          e_cost = st.number_input("Unit Cost (GH₵)", min_value=0.0, step=0.5, value=float(selected_row["unit_cost"]), format="%.2f", key=f"{key_prefix}_u")

        if st.form_submit_button("💾 Save Changes", type="primary", use_container_width=True):
          payload = {
              "date": e_date,
              "customer_name": e_cust.strip().upper(),
              "batch_number": e_batch.strip().upper(),
              "product_name": e_prod,
              "weight_g": float(e_weight),
              "quantity": int(e_qty),
              "unit_cost": float(e_cost),
              "entered_by": e_staff.strip(),
          }
          update_dispatch_record(table_name, rec_id, payload)
          st.success(f"Record #{rec_id} successfully updated!")
          st.rerun()
    else:
      if st.button("🗑 Permanently Delete This Record", type="primary", use_container_width=True):
        delete_record(table_name, rec_id)
        st.warning(f"Record #{rec_id} permanently deleted.")
        st.rerun()


tab_entry, tab_local, tab_export, tab_protocol, tab_qa, tab_analytics = st.tabs(
    [
        "📝 New Sales Entry",
        "📦 Local Sales Table",
        "🚢 Exports Table",
        "🏛️ Protocol Table",
        "🔍 QA Desk",
        "📊 Analytics & Trends",
    ]
)

with tab_entry:
  st.subheader("Sales Details (Packaging / Entry)")
  
  if "current_weight" not in st.session_state:
    st.session_state.current_weight = 50.0
  if "current_cost" not in st.session_state:
    st.session_state.current_cost = 16.0

  def update_product_defaults():
    prod = st.session_state.get("product_selectbox", "-- Select Product --")
    destination_val = st.session_state.get("target_destination", "Local Sales")
    
    if prod != "-- Select Product --" and not CATALOG_DF.empty:
      match_row = CATALOG_DF[CATALOG_DF["Product Name"] == prod]
      if not match_row.empty:
        raw_w = match_row["Weight"].values[0]
        if "almond" in prod.lower() and "27g" not in prod.lower():
          st.session_state.current_weight = 80.0
        else:
          st.session_state.current_weight = parse_weight_to_grams(raw_w, prod)
        
        if destination_val == "Exports":
          st.session_state.current_cost = float(match_row["Export Price (GHS)"].fillna(50.0).values[0])
        else:
          st.session_state.current_cost = float(match_row["Local Price (GHS)"].fillna(25.0).values[0])

  c1, c2, c3 = st.columns([2, 2, 2])
  with c1:
    destination = st.selectbox("Target Section", ["Local Sales", "Exports", "Protocol"], key="target_destination", on_change=update_product_defaults)
  with c2:
    entry_date = st.date_input("Date", value=date.today())
  with c3:
    entered_by = st.text_input("Entered By (Staff Name)", value=st.session_state.get("staff_user", ""))

  c4, c5, c_batch = st.columns([3, 3, 2])
  with c4:
    customer_name = st.text_input("Customer Name", placeholder="e.g. Tomreik Hotel")
  with c5:
    product_name = st.selectbox("Product Name", PRODUCT_CATALOG, key="product_selectbox", on_change=update_product_defaults)
  with c_batch:
    batch_number = st.text_input("Batch Number", placeholder="e.g. BATCH-2026-001")

  c6, c7, c8 = st.columns(3)
  with c6:
    unit_weight_g = st.number_input("Weight per Unit (g)", value=float(st.session_state.current_weight), format="%.2f", key="current_weight")
  with c7:
    quantity = st.number_input("Quantity", min_value=1, value=1)
  with c8:
    unit_cost = st.number_input("Unit Cost (GH₵)", value=float(st.session_state.current_cost), format="%.2f", key="current_cost")

  calc_total_g = quantity * unit_weight_g
  calc_total_mt = calc_total_g / 1_000_000.0
  calc_total_amt = quantity * unit_cost

  st.markdown(
      f"""
      <div style="background: rgba(224, 169, 109, 0.15); padding: 14px 20px; border-radius: 12px; margin: 15px 0; border-left: 4px solid #f4a261; color: #ffffff;">
          <strong>Live Computation Preview:</strong><br>
          Total Weight: <b>{calc_total_g:,.2f} g</b> &nbsp;|&nbsp; 
          Total Weight in Metric Tonnes: <b>{calc_total_mt:,.6f} MT</b> &nbsp;|&nbsp; 
          Total Amount (GH₵): <b>GH₵ {calc_total_amt:,.2f}</b>
      </div>
      """,
      unsafe_allow_html=True,
  )

  if st.button("Record Sale", type="primary", use_container_width=True):
    if not customer_name.strip() or product_name == "-- Select Product --":
      st.error("Please provide a Customer Name and select a Product.")
    else:
      payload = {
          "date": entry_date,
          "customer_name": customer_name.strip().upper(),
          "batch_number": batch_number.strip().upper() if batch_number.strip() else "N/A",
          "product_name": product_name,
          "weight_g": float(unit_weight_g),
          "quantity": int(quantity),
          "unit_cost": float(unit_cost),
          "entered_by": entered_by.strip(),
      }
      try:
        insert_dispatch(TABLE_MAP[destination], payload)
        st.success(f"Sale for '{customer_name}' recorded successfully!")
      except Exception as e:
        st.error(f"Database error: {e}")


def render_sales_table_view(table_key, title):
  st.subheader(f"{title} Records")
  try:
    df = fetch_table(table_key)
    if df.empty:
      st.info(f"No records found in {table_key} yet.")
      return

    df["date"] = pd.to_datetime(df["date"]).dt.date
    if "batch_number" not in df.columns:
      df["batch_number"] = "-"

    render_table_edit_delete_bar(table_key, key_prefix=f"{table_key}_panel")

    f_col1, f_col2 = st.columns([4, 4])
    with f_col1:
      search_query = st.text_input(
          "🔍 Search Client / Batch / Product / Staff",
          key=f"search_{table_key}",
          placeholder="Type to filter...",
      )
    with f_col2:
      min_date = df["date"].min()
      max_date = df["date"].max()
      date_range = st.date_input(
          "Filter Date Range",
          value=(min_date, max_date),
          min_value=min_date,
          max_value=max_date,
          key=f"date_{table_key}",
      )

    filtered_df = df.copy()
    if search_query.strip():
      q = search_query.strip().lower()
      filtered_df = filtered_df[
          filtered_df["customer_name"].str.lower().str.contains(q, na=False)
          | filtered_df["batch_number"].astype(str).str.lower().str.contains(q, na=False)
          | filtered_df["product_name"].str.lower().str.contains(q, na=False)
          | filtered_df["entered_by"].str.lower().str.contains(q, na=False)
      ]
    if isinstance(date_range, (tuple, list)) and len(date_range) == 2:
      start_d, end_d = date_range
      filtered_df = filtered_df[
          (filtered_df["date"] >= start_d) & (filtered_df["date"] <= end_d)
      ]

    tot_mt = filtered_df["total_weight_mt"].sum()
    tot_ghc = filtered_df["total_amount"].sum()
    tot_usd = (
        (filtered_df["total_amount"] / filtered_df["exchange_rate"]).sum()
        if "exchange_rate" in filtered_df.columns
        else 0.0
    )

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Orders / Rows", len(filtered_df))
    k2.metric("Total Weight in Metric Tonnes", f"{tot_mt:,.6f} MT")
    k3.metric("Total Amount (GH₵)", f"GH₵ {tot_ghc:,.2f}")
    k4.metric("Total Amount (USD)", f"${tot_usd:,.2f}")

    display_merged_excel_table(filtered_df)

    export_df = prepare_export_df(filtered_df)
    e_col1, e_col2, _ = st.columns([2, 2, 6])
    with e_col1:
      csv_data = export_df.to_csv(index=False).encode("utf-8")
      st.download_button(
          label="📥 Download CSV",
          data=csv_data,
          file_name=f"{table_key}_{date.today()}.csv",
          mime="text/csv",
          use_container_width=True,
      )
    with e_col2:
      try:
        excel_data = to_excel_bytes(filtered_df)
        st.download_button(
            label="📊 Download Excel (.xlsx)",
            data=excel_data,
            file_name=f"{table_key}_{date.today()}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
      except Exception as ex:
        st.error(f"Error generating Excel: {ex}")

  except Exception as e:
    st.error(f"Failed to load table: {e}")


with tab_local:
  render_sales_table_view("local_sales", "Local Sales")

with tab_export:
  render_sales_table_view("export_sales", "Exports")

with tab_protocol:
  render_sales_table_view("protocol_sales", "Protocol")

with tab_qa:
  st.subheader("🔍 QA Desk")
  st.caption("Apply monthly exchange rates or edit rates box-by-box to calculate USD valuations automatically.")

  q_sec, q_mode = st.columns([3, 3])
  with q_sec:
    qa_section = st.selectbox(
        "Select Target Section",
        ["Local Sales", "Exports", "Protocol"],
        key="qa_desk_sec",
    )
  with q_mode:
    qa_mode = st.radio(
        "Update Method",
        ["Apply by Entire Month (Bulk)", "Edit Box-by-Box (Table Editor)"],
        horizontal=True,
    )

  qa_table_name = TABLE_MAP[qa_section]
  raw_qa_df = fetch_table(qa_table_name)

  if not raw_qa_df.empty:
    raw_qa_df["date"] = pd.to_datetime(raw_qa_df["date"]).dt.date
    if "batch_number" not in raw_qa_df.columns:
      raw_qa_df["batch_number"] = "-"

    if qa_mode == "Apply by Entire Month (Bulk)":
      st.markdown("##### 📅 Set Monthly Exchange Rate")
      b_c1, b_c2, b_c3, b_c4 = st.columns([2, 2, 2, 3])
      with b_c1:
        qa_year = st.selectbox("Year", [2025, 2026, 2027], index=1)
      with b_c2:
        month_names = [
            "January", "February", "March", "April", "May", "June",
            "July", "August", "September", "October", "November", "December",
        ]
        qa_month_name = st.selectbox("Month", month_names, index=date.today().month - 1)
        qa_month = month_names.index(qa_month_name) + 1
      with b_c3:
        qa_bulk_rate = st.number_input("Dollar Rate (GH₵/$)", min_value=0.01, step=0.05, value=11.20, format="%.4f")
      with b_c4:
        st.write("")
        st.write("")
        if st.button(f"Apply {qa_bulk_rate:.2f} to All {qa_month_name} {qa_year} Records", type="primary"):
          apply_monthly_rate(qa_table_name, qa_year, qa_month, qa_bulk_rate)
          st.success(f"Applied exchange rate of {qa_bulk_rate:.4f} to all {qa_month_name} {qa_year} entries!")
          st.rerun()

    else:
      st.markdown("##### ✏ Double-Click the Exchange Rate Box to Edit Directly")
      editable_df = raw_qa_df[[
          "id", "date", "customer_name", "batch_number", "product_name",
          "quantity", "unit_cost", "total_amount", "exchange_rate",
      ]].copy()

      edited_result = st.data_editor(
          editable_df,
          disabled=[
              "id", "date", "customer_name", "batch_number", "product_name",
              "quantity", "unit_cost", "total_amount",
          ],
          column_config={
              "exchange_rate": st.column_config.NumberColumn(
                  "Exchange Rate (GH₵/$)", min_value=0.01, format="%.4f"
              )
          },
          hide_index=True,
          use_container_width=True,
          key="qa_box_editor",
      )

      if st.button("Save Box Changes to Database", type="primary"):
        for _, row in edited_result.iterrows():
          update_row_rate(qa_table_name, int(row["id"]), float(row["exchange_rate"]))
        st.success("All box-by-box rate changes saved to database!")
        st.rerun()

    render_table_edit_delete_bar(qa_table_name, key_prefix=f"qa_direct_{qa_table_name}")

    st.markdown("---")
    st.markdown("### 📋 Sales Data")

    tot_ledger_mt = raw_qa_df["total_weight_mt"].sum()
    tot_ledger_ghc = raw_qa_df["total_amount"].sum()
    tot_ledger_usd = (
        (raw_qa_df["total_amount"] / raw_qa_df["exchange_rate"]).sum()
        if "exchange_rate" in raw_qa_df.columns
        else 0.0
    )

    qk1, qk2, qk3, qk4 = st.columns(4)
    qk1.metric("Orders / Rows", len(raw_qa_df))
    qk2.metric("Total Weight in Metric Tonnes", f"{tot_ledger_mt:,.6f} MT")
    qk3.metric("Total Amount (GH₵)", f"GH₵ {tot_ledger_ghc:,.2f}")
    qk4.metric("Total Amount (USD)", f"${tot_ledger_usd:,.2f}")

    display_merged_excel_table(raw_qa_df)

    export_qa_df = prepare_export_df(raw_qa_df)
    d1, d2, _ = st.columns([2, 2, 6])
    with d1:
      st.download_button(
          label="📥 Export QA Sales Data CSV",
          data=export_qa_df.to_csv(index=False).encode("utf-8"),
          file_name=f"QA_Sales_Data_{qa_table_name}_{date.today()}.csv",
          mime="text/csv",
          use_container_width=True,
      )
    with d2:
      try:
        st.download_button(
            label="📊 Export QA Sales Data Excel (.xlsx)",
            data=to_excel_bytes(raw_qa_df),
            file_name=f"QA_Sales_Data_{qa_table_name}_{date.today()}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
      except Exception as ex:
        st.error(f"Error generating Excel: {ex}")
  else:
    st.info(f"No records logged in {qa_section} yet.")

with tab_analytics:
  st.subheader("📊 Sales & Production Performance")
  try:
    df_l = fetch_table("local_sales")
    df_e = fetch_table("export_sales")
    df_p = fetch_table("protocol_sales")

    for df, channel_name, table_label in [
        (df_l, "Local Sales", "Local Sales Table"),
        (df_e, "Exports", "Exports Table"),
        (df_p, "Protocol", "Protocol Table")
    ]:
      df["Channel"] = channel_name
      df["Source Table"] = table_label
      if "total_weight_mt" not in df.columns:
        df["total_weight_mt"] = (df["quantity"] * df["weight_g"] / 1_000_000.0) if not df.empty else 0.0
      if "total_amount" not in df.columns:
        df["total_amount"] = (df["quantity"] * df["unit_cost"]) if not df.empty else 0.0
      if "exchange_rate" not in df.columns:
        df["exchange_rate"] = 1.0000

    all_df = pd.concat([df_l, df_e, df_p], ignore_index=True)

    if not all_df.empty:
      all_df["total_weight_mt"] = pd.to_numeric(all_df["total_weight_mt"], errors="coerce").fillna(0.0)
      all_df["total_amount"] = pd.to_numeric(all_df["total_amount"], errors="coerce").fillna(0.0)

      local_mt = df_l["total_weight_mt"].sum() if not df_l.empty else 0.0
      local_ghc = df_l["total_amount"].sum() if not df_l.empty else 0.0

      protocol_mt = df_p["total_weight_mt"].sum() if not df_p.empty else 0.0
      protocol_ghc = df_p["total_amount"].sum() if not df_p.empty else 0.0

      exports_mt = df_e["total_weight_mt"].sum() if not df_e.empty else 0.0
      exports_ghc = df_e["total_amount"].sum() if not df_e.empty else 0.0

      tot_grand_mt = local_mt + protocol_mt + exports_mt
      tot_grand_ghc = local_ghc + protocol_ghc + exports_ghc
      tot_grand_usd = (all_df["total_amount"] / all_df["exchange_rate"]).sum() if not all_df.empty else 0.0

      proto_weight_share = (protocol_mt / tot_grand_mt * 100.0) if tot_grand_mt > 0 else 0.0
      proto_value_share = (protocol_ghc / tot_grand_ghc * 100.0) if tot_grand_ghc > 0 else 0.0

      a_col1, a_col2, a_col3, a_col4 = st.columns(4)
      a_col1.metric("Total Orders", f"{len(all_df)}")
      a_col2.metric("Total Weight (MT)", f"{tot_grand_mt:,.6f} MT")
      a_col3.metric("Total Money (GH₵)", f"GH₵ {tot_grand_ghc:,.2f}")
      a_col4.metric("Total Money (USD)", f"${tot_grand_usd:,.2f}")

      st.markdown("---")
      st.markdown("### ⚖️ Weight Comparison: Domestic vs Exports")

      domestic_mt = local_mt + protocol_mt
      domestic_pct = (domestic_mt / tot_grand_mt * 100.0) if tot_grand_mt > 0 else 0.0
      exports_pct = (exports_mt / tot_grand_mt * 100.0) if tot_grand_mt > 0 else 0.0

      c_m1, c_m2 = st.columns(2)
      with c_m1:
        st.metric(
            label="Domestic Weight (Local + Protocol)",
            value=f"{domestic_mt:,.6f} MT",
            delta=f"{domestic_pct:.2f}% of Total Weight",
        )
      with c_m2:
        st.metric(
            label="Exports Weight",
            value=f"{exports_mt:,.6f} MT",
            delta=f"{exports_pct:.2f}% of Total Weight",
        )

      comp_df = pd.DataFrame({
          "Market": ["Domestic", "Exports"],
          "Weight (MT)": [domestic_mt, exports_mt],
          "Percentage": [domestic_pct, exports_pct],
      })

      c_chart1, c_chart2 = st.columns(2)
      with c_chart1:
        st.markdown("#### Percentage of Weight: Domestic vs Exports")
        fig_donut = px.pie(
            comp_df,
            names="Market",
            values="Weight (MT)",
            hole=0.45,
            color="Market",
            color_discrete_map={"Domestic": "#1E88E5", "Exports": "#FFFFFF"},
        )
        fig_donut.update_traces(marker=dict(line=dict(color="#000000", width=2)))
        fig_donut.update_layout(
            margin=dict(t=20, b=20, l=10, r=10),
            showlegend=False,
            height=320,
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="rgba(0,0,0,0)",
            font=dict(color="#ffffff"),
        )
        st.plotly_chart(fig_donut, use_container_width=True)

      with c_chart2:
        st.markdown("#### Total Weight in Metric Tonnes: Domestic vs Exports")
        fig_bar_comp = px.bar(
            comp_df,
            x="Market",
            y="Weight (MT)",
            text="Weight (MT)",
            color="Market",
            color_discrete_map={"Domestic": "#1E88E5", "Exports": "#FFFFFF"},
        )
        fig_bar_comp.update_traces(texttemplate="%{text:.4f} MT", textposition="outside", marker_line_color="#000000", marker_line_width=1.5)
        fig_bar_comp.update_layout(
            xaxis=dict(categoryorder="array", categoryarray=["Domestic", "Exports"]),
            yaxis_title="Metric Tonnes (MT)",
            xaxis_title="",
            margin=dict(t=20, b=20, l=10, r=10),
            showlegend=False,
            height=320,
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="rgba(0,0,0,0)",
            font=dict(color="#ffffff"),
        )
        st.plotly_chart(fig_bar_comp, use_container_width=True)

      st.markdown("---")
      st.markdown("### 🍫 Most Sold Products by Each Channel")
      
      prod_cols = st.columns(3)
      channel_data_map = [
          ("Top Products: Local Sales", df_l, "Local Sales Table only"),
          ("Top Products: Exports", df_e, "Exports Table only"),
          ("Top Products: Protocol", df_p, "Protocol Table only")
      ]

      for i, (title, channel_df, src_label) in enumerate(channel_data_map):
        with prod_cols[i]:
          st.markdown(f"#### {title}")
          st.caption(f"📌 Source: {src_label}")
          if not channel_df.empty and "product_name" in channel_df.columns and "quantity" in channel_df.columns:
            top_prod = channel_df.groupby("product_name")["quantity"].sum().reset_index()
            top_prod = top_prod.sort_values(by="quantity", ascending=False).head(7)
            
            bar_color = "#1E88E5" if "Local Sales" in title else ("#FFFFFF" if "Exports" in title else "#FFD700")

            fig_prod = px.bar(
                top_prod,
                x="product_name",
                y="quantity",
                text="quantity",
                color_discrete_sequence=[bar_color]
            )
            fig_prod.update_traces(texttemplate="%{text}", textposition="outside", marker_color=bar_color, marker_line_color="#000000", marker_line_width=1.5)
            fig_prod.update_layout(
                xaxis_title="",
                yaxis_title="Units Sold",
                margin=dict(t=20, b=40, l=10, r=10),
                height=280,
                paper_bgcolor="rgba(0,0,0,0)",
                plot_bgcolor="rgba(0,0,0,0)",
                font=dict(color="#ffffff"),
                xaxis=dict(tickangle=-45)
            )
            st.plotly_chart(fig_prod, use_container_width=True)
          else:
            st.info("No records in table.")

      st.markdown("---")
      st.markdown("### 📊 Orders and Money by Channel")
      
      channel_summary_df = pd.DataFrame({
          "Channel": ["Local Sales", "Exports", "Protocol"],
          "Orders": [len(df_l), len(df_e), len(df_p)],
          "Amount": [local_ghc, exports_ghc, protocol_ghc]
      })

      ch_col1, ch_col2 = st.columns(2)
      with ch_col1:
        st.markdown("#### Number of Orders by Channel")
        st.caption("📌 Source: Total order rows in Local Sales Table, Protocol Table, and Exports Table separately")
        fig_ch_orders = px.bar(
            channel_summary_df,
            x="Channel",
            y="Orders",
            text="Orders",
            color="Channel",
            color_discrete_map={"Local Sales": "#1E88E5", "Exports": "#FFFFFF", "Protocol": "#FFD700"}
        )
        fig_ch_orders.update_traces(texttemplate="%{text}", textposition="outside", marker_line_color="#000000", marker_line_width=1.5)
        fig_ch_orders.update_layout(
            xaxis=dict(categoryorder="array", categoryarray=["Local Sales", "Exports", "Protocol"]),
            xaxis_title="",
            yaxis_title="Orders",
            margin=dict(t=20, b=20, l=10, r=10),
            showlegend=False,
            height=300,
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="rgba(0,0,0,0)",
            font=dict(color="#ffffff")
        )
        st.plotly_chart(fig_ch_orders, use_container_width=True)

      with ch_col2:
        st.markdown("#### Total Money (GH₵) by Channel")
        st.caption("📌 Source: Sum of money recorded in Local Sales Table, Protocol Table, and Exports Table separately")
        fig_ch_money = px.bar(
            channel_summary_df,
            x="Channel",
            y="Amount",
            text="Amount",
            color="Channel",
            color_discrete_map={"Local Sales": "#1E88E5", "Exports": "#FFFFFF", "Protocol": "#FFD700"}
        )
        fig_ch_money.update_traces(texttemplate="GH₵ %{text:,.2f}", textposition="outside", marker_line_color="#000000", marker_line_width=1.5)
        fig_ch_money.update_layout(
            xaxis=dict(categoryorder="array", categoryarray=["Local Sales", "Exports", "Protocol"]),
            xaxis_title="",
            yaxis_title="Amount (GH₵)",
            margin=dict(t=20, b=20, l=10, r=10),
            showlegend=False,
            height=300,
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="rgba(0,0,0,0)",
            font=dict(color="#ffffff")
        )
        st.plotly_chart(fig_ch_money, use_container_width=True)

      st.markdown("---")
      st.markdown("### 🏛️ Protocol Tracker")
      pr_col1, pr_col2, pr_col3 = st.columns(3)
      with pr_col1:
        st.metric(
            label="Protocol Weight",
            value=f"{protocol_mt:,.6f} MT",
            delta=f"{proto_weight_share:.2f}% of All Weight",
        )
      with pr_col2:
        st.metric(
            label="Protocol Money Value",
            value=f"GH₵ {protocol_ghc:,.2f}",
            delta=f"{proto_value_share:.2f}% of All Money",
        )
      with pr_col3:
        st.metric(
            label="Protocol Orders",
            value=f"{len(df_p)} orders",
            delta=f"{(len(df_p) / len(all_df) * 100 if len(all_df) > 0 else 0):.1f}% of total",
        )

    else:
      st.info("No orders found across the tables to display graphs.")
  except Exception as e:
    st.error(f"Failed to generate graphs: {e}")
